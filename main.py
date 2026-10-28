
# -*- coding: utf-8 -*-
"""
Главный скрипт поиска цен на 24lek.ru.

Ищет цены:
1. В ближайших физических аптеках.
2. В фиксированном списке интернет-аптек.
3. Повторно ищет ненайденные цены физических аптек
   с фильтром по району.

В отчёте:
- лист «Отчёт» — цены и статусы;
- лист «Детали» — карточки товаров и источники.

Запуск:
    python main.py
"""

import sys
import time

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

import config
from normalizer import normalize_name
from product_matcher import match_product
from geocoder import geocode_address, GeocodeError
from pharmacies import (
    load_pharmacies,
    pharmacies_in_radius,
    pharmacy_header,
)
from lek24_parser import search_drug_prices, _make_session


INTERNET_PHARMACY_IDS = [
    1105,  # Аптека.ру
    731,   # Аптека Легко
    897,   # Губернские аптеки
    1151,  # Фармэконом
    344,   # Нейрон
    813,   # ТаблеткиТут
]


DISTRICT_RAON_IDS = {
    "железнодорожный": 1,
    "кировский": 2,
    "ленинский": 3,
    "октябрьский": 4,
    "свердловский": 5,
    "советский": 6,
    "центральный": 7,
}


def normalize_id(value):
    """Приводит ID аптеки к целому числу, если возможно."""
    try:
        return int(value)
    except (TypeError, ValueError):
        return value


def normalize_district(value):
    """Нормализует название района для сопоставления."""
    if value is None:
        return ""

    district = " ".join(str(value).strip().lower().split())
    district = district.removesuffix(" район").strip()

    return district


def get_raon_id(pharmacy):
    """Возвращает ID района аптеки либо None."""
    district = normalize_district(pharmacy.get("district"))
    return DISTRICT_RAON_IDS.get(district)


def read_products(path: str) -> list:
    """Читает названия препаратов из первого непустого столбца."""
    wb = openpyxl.load_workbook(
        path,
        read_only=True,
        data_only=True,
    )
    ws = wb.active
    names = []

    try:
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell is None:
                    continue

                text = str(cell).strip()
                if text:
                    names.append(text)
                break
    finally:
        wb.close()

    return names


def get_record(price_map, pharmacy_id):
    """Возвращает запись о найденной карточке."""
    record = price_map.get(normalize_id(pharmacy_id))

    if isinstance(record, dict):
        return record

    # Совместимость с прежним форматом: ID -> цена.
    if isinstance(record, (int, float)):
        return {
            "price": record,
            "name": "",
            "source": "24lek.ru",
            "status": "НАЙДЕНО",
            "reason": "Цена найдена; сведения о карточке не сохранены.",
        }

    return None


def get_price(price_map, pharmacy_id):
    """Возвращает цену из записи."""
    record = get_record(price_map, pharmacy_id)

    if not record:
        return None

    return record.get("price")


def get_pharmacy_status(price_map, pharmacy_id):
    """
    Возвращает статус для отчёта.

    Пока парсер не сообщает достоверно, что поиск полностью
    завершён, отсутствие цены не считается доказанным отсутствием
    товара.
    """
    record = get_record(price_map, pharmacy_id)

    if record and record.get("price") is not None:
        return "НАЙДЕНО"

    if record and record.get("status"):
        return record["status"]

    return "НЕ ПОДТВЕРЖДЕНО"


def build_report(
    raw_products: list,
    nearby_pharmacies: list,
    internet_pharmacies: list,
    results: dict,
    out_path: str,
):
    """Создаёт Excel-отчёт с ценами, статусами и деталями карточек."""
    wb = openpyxl.Workbook()

    # ---------------------------------------------------------
    # Лист 1. Сводный отчёт
    # ---------------------------------------------------------
    ws = wb.active
    ws.title = "Отчёт"

    all_pharmacies = nearby_pharmacies + internet_pharmacies

    headers = (
        ["Препарат (из списка)"]
        + [
            pharmacy_header(ph)
            for ph in nearby_pharmacies
        ]
        + [
            f"ИНТЕРНЕТ | {ph['name']} | {ph['address']}"
            for ph in internet_pharmacies
        ]
    )

    ws.append(headers)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(
        start_color="2F5496",
        end_color="2F5496",
        fill_type="solid",
    )
    internet_fill = PatternFill(
        start_color="548235",
        end_color="548235",
        fill_type="solid",
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.alignment = Alignment(
            wrap_text=True,
            vertical="center",
            horizontal="center",
        )

        if col_idx > 1 + len(nearby_pharmacies):
            cell.fill = internet_fill
        else:
            cell.fill = header_fill

    for raw_name in raw_products:
        price_map = results.get(raw_name, {})
        row = [raw_name]

        for ph in all_pharmacies:
            aid = normalize_id(ph["id"])
            price = get_price(price_map, aid)
            status = get_pharmacy_status(price_map, aid)

            # Цена остаётся числом, чтобы её можно было использовать
            # в формулах Excel. Если цены нет, показываем статус.
            if price is not None:
                row.append(price)
            else:
                row.append(status)

        ws.append(row)

    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 48

    for col_idx in range(2, len(headers) + 1):
        ws.column_dimensions[
            get_column_letter(col_idx)
        ].width = 28

    ws.row_dimensions[1].height = 75

    for row in ws.iter_rows(min_row=2, min_col=2):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0.00'

    # ---------------------------------------------------------
    # Лист 2. Подробности по каждой паре препарат / аптека
    # ---------------------------------------------------------
    details = wb.create_sheet("Детали")

    detail_headers = [
        "Препарат из списка",
        "Аптека",
        "ID аптеки",
        "Тип аптеки",
        "Статус",
        "Цена, руб.",
        "Название карточки на 24lek.ru",
        "Источник",
        "Комментарий",
    ]

    details.append(detail_headers)

    for col_idx in range(1, len(detail_headers) + 1):
        cell = details.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(
            wrap_text=True,
            vertical="center",
            horizontal="center",
        )

    internet_ids = {
        normalize_id(ph["id"])
        for ph in internet_pharmacies
    }

    for raw_name in raw_products:
        price_map = results.get(raw_name, {})

        for ph in all_pharmacies:
            aid = normalize_id(ph["id"])
            record = get_record(price_map, aid)
            price = get_price(price_map, aid)
            status = get_pharmacy_status(price_map, aid)

            if aid in internet_ids:
                pharmacy_type = "Интернет-аптека"
            else:
                pharmacy_type = "Физическая аптека"

            if record:
                card_name = record.get("name", "")
                source = record.get("source", "24lek.ru")
                reason = record.get("reason", "")
            else:
                card_name = ""
                source = "24lek.ru"
                reason = (
                    "Подходящая карточка не сохранена. "
                    "Полное отсутствие товара не подтверждено."
                )

            details.append(
                [
                    raw_name,
                    ph.get("name", ""),
                    aid,
                    pharmacy_type,
                    status,
                    price if price is not None else "",
                    card_name,
                    source,
                    reason,
                ]
            )

    details.freeze_panes = "A2"
    details.auto_filter.ref = details.dimensions
    details.row_dimensions[1].height = 40

    detail_widths = {
        "A": 42,
        "B": 34,
        "C": 12,
        "D": 20,
        "E": 24,
        "F": 14,
        "G": 60,
        "H": 36,
        "I": 65,
    }

    for column, width in detail_widths.items():
        details.column_dimensions[column].width = width

    for row in details.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

        price_cell = row[5]
        if isinstance(price_cell.value, (int, float)):
            price_cell.number_format = '#,##0.00'

    wb.save(out_path)


def update_price_map(
    price_map,
    items,
    allowed_ids,
    source_product,
):
    """
    Добавляет подходящие карточки.

    Сохраняет минимальную найденную цену на аптеку вместе
    с названием карточки и доступным источником.
    """
    for item in items:
        aid = normalize_id(item.get("apteka_id"))

        if aid not in allowed_ids:
            continue

        price = item.get("price")

        if price is None:
            continue

        card_name = str(item.get("name", "")).strip()
        match = match_product(source_product, card_name)

        if not match["matches"]:
            print(
                f"      Карточка отклонена [аптека ID {aid}]: "
                f"{card_name!r}. Причина: {match['reason']}"
            )
            continue

        # Парсер в текущей версии не передаёт прямую ссылку
        # на карточку, поэтому сохраняем источник поиска.
        source = item.get("source") or "24lek.ru"

        new_record = {
            "price": float(price),
            "name": card_name,
            "source": source,
            "status": "НАЙДЕНО",
            "reason": "Найдено совпадение по характеристикам препарата.",
        }

        old_record = get_record(price_map, aid)

        if (
            old_record is None
            or old_record.get("price") is None
            or new_record["price"] < old_record["price"]
        ):
            price_map[aid] = new_record


def main():
    print("=" * 70)
    print(" Поиск цен на лекарства на 24lek.ru")
    print(" Ближайшие аптеки + интернет-аптеки")
    print("=" * 70)

    address = input(
        "\nВведите адрес (например: Свободный 60): "
    ).strip()

    if not address:
        print("Адрес не введён. Завершение.")
        sys.exit(1)

    print("\n[1/5] Определяем координаты адреса через Dadata...")

    try:
        lat, lon, full_address = geocode_address(address)
    except GeocodeError as e:
        print(f"ОШИБКА геокодирования: {e}")
        sys.exit(1)

    print(f"   Адрес распознан как: {full_address}")
    print(f"   Координаты: {lat:.6f}, {lon:.6f}")

    print(
        f"\n[2/5] Загружаем {config.APTEKI_FILE} "
        f"и определяем аптеки..."
    )

    try:
        pharmacies = load_pharmacies()
    except FileNotFoundError:
        print(
            f"ОШИБКА: файл {config.APTEKI_FILE} "
            "не найден рядом со скриптом."
        )
        sys.exit(1)

    internet_ids = {
        normalize_id(aid)
        for aid in INTERNET_PHARMACY_IDS
    }

    pharmacies_by_id = {
        normalize_id(ph["id"]): ph
        for ph in pharmacies
    }

    internet_pharmacies = []

    for aid in INTERNET_PHARMACY_IDS:
        ph = pharmacies_by_id.get(normalize_id(aid))

        if ph is None:
            print(
                f"ПРЕДУПРЕЖДЕНИЕ: интернет-аптека с ID {aid} "
                "не найдена в apteki.json."
            )
            continue

        internet_pharmacies.append(ph)

    nearby_all = pharmacies_in_radius(
        pharmacies,
        lat,
        lon,
    )

    nearby = [
        ph for ph in nearby_all
        if normalize_id(ph["id"]) not in internet_ids
    ]

    print(f"   Ближайших физических аптек: {len(nearby)}")
    print(
        f"   Интернет-аптек из списка найдено: "
        f"{len(internet_pharmacies)}"
    )

    for ph in nearby:
        print(
            f"     - {ph['name']} — {ph['address']} "
            f"({ph['distance_m']} м) [id={ph['id']}]"
        )

    print("\n   Интернет-аптеки:")
    for ph in internet_pharmacies:
        print(
            f"     - {ph['name']} — {ph['address']} "
            f"[id={ph['id']}]"
        )

    needed_ids = {
        normalize_id(ph["id"])
        for ph in nearby + internet_pharmacies
    }

    if not needed_ids:
        print(
            "Не найдено аптек для поиска цен. "
            "Проверьте apteki.json."
        )
        sys.exit(1)

    print(
        f"\n[3/5] Загружаем список препаратов "
        f"из {config.PRODUCTS_FILE}..."
    )

    try:
        raw_products = read_products(config.PRODUCTS_FILE)
    except FileNotFoundError:
        print(
            f"ОШИБКА: файл {config.PRODUCTS_FILE} "
            "не найден рядом со скриптом."
        )
        sys.exit(1)

    print(f"   Загружено препаратов: {len(raw_products)}")

    print(
        "\n[4/5] Ищем цены для ближайших и интернет-аптек..."
    )

    physical_ids = {
        normalize_id(ph["id"])
        for ph in nearby
    }

    pharmacies_by_raon = {}

    for ph in nearby:
        aid = normalize_id(ph["id"])
        raon_id = get_raon_id(ph)

        if raon_id is None:
            print(
                f"[Районный поиск] Неизвестен район аптеки "
                f"{ph['name']} [id={aid}]: "
                f"{ph.get('district')!r}. "
                "Резервный поиск для неё пропускается."
            )
            continue

        pharmacies_by_raon.setdefault(
            raon_id, set()
        ).add(aid)

    session = _make_session()
    results = {}

    try:
        for idx, raw_name in enumerate(raw_products, start=1):
            norm = normalize_name(raw_name)

            if not norm:
                print(
                    f"   [{idx}/{len(raw_products)}] "
                    f"Пропуск: не удалось нормализовать "
                    f"{raw_name!r}"
                )
                results[raw_name] = {}
                continue

            print(
                f"\n   [{idx}/{len(raw_products)}] "
                f"{raw_name!r} -> запрос: '{norm}'"
            )

            try:
                items = search_drug_prices(
                    norm,
                    needed_apteka_ids=needed_ids,
                    session=session,
                    raon_id=0,
                )
            except Exception as e:
                print(
                    f"      ОШИБКА основного поиска: {e}"
                )
                items = []

            price_map = {}
            update_price_map(
                price_map,
                items,
                needed_ids,
                raw_name,
            )

            missing_physical_ids = (
                physical_ids - set(price_map)
            )

            if missing_physical_ids:
                print(
                    f"      Районный поиск: цены пока не найдены "
                    f"для {len(missing_physical_ids)} "
                    f"физических аптек."
                )

                for raon_id, district_ids in sorted(
                    pharmacies_by_raon.items()
                ):
                    target_ids = (
                        missing_physical_ids & district_ids
                    )

                    if not target_ids:
                        continue

                    district_name = next(
                        (
                            str(ph.get("district", "")).strip()
                            for ph in nearby
                            if normalize_id(ph["id"])
                            in target_ids
                        ),
                        f"raon={raon_id}",
                    )

                    print(
                        f"      Ищем в районе "
                        f"{district_name}: "
                        f"нужные ID {sorted(target_ids)}"
                    )

                    try:
                        district_items = search_drug_prices(
                            norm,
                            needed_apteka_ids=target_ids,
                            session=session,
                            raon_id=raon_id,
                        )
                    except Exception as e:
                        print(
                            f"      ОШИБКА районного поиска "
                            f"raon={raon_id}: {e}"
                        )
                        continue

                    update_price_map(
                        price_map,
                        district_items,
                        target_ids,
                        raw_name,
                    )

                    missing_physical_ids = (
                        physical_ids - set(price_map)
                    )

            physical_found = sum(
                1
                for ph in nearby
                if get_price(
                    price_map,
                    ph["id"],
                ) is not None
            )

            internet_found = sum(
                1
                for ph in internet_pharmacies
                if get_price(
                    price_map,
                    ph["id"],
                ) is not None
            )

            print(
                f"      Физических аптек с ценой: "
                f"{physical_found}/{len(nearby)}"
            )
            print(
                f"      Интернет-аптек с ценой: "
                f"{internet_found}/{len(internet_pharmacies)}"
            )

            results[raw_name] = price_map

            if config.REQUEST_DELAY:
                time.sleep(config.REQUEST_DELAY)

    finally:
        session.close()

    print(f"\n[5/5] Формируем отчёт {config.REPORT_FILE}...")

    build_report(
        raw_products,
        nearby,
        internet_pharmacies,
        results,
        config.REPORT_FILE,
    )

    print(f"\nГотово! Отчёт сохранён: {config.REPORT_FILE}")


if __name__ == "__main__":
    main()
